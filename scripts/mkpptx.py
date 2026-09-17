#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
mkpptx.py —— 把已定稿的贝莱德风格 PDF 逐页封装为 PPTX。

策略（与 brand-pdf-deck-replica 同源，但页面尺寸自动探测，适配 A4 横向）：
  · 幻灯片尺寸 = PDF 首页尺寸（本例 A4 横向 841.92×594.96pt），1:1 不缩放
  · 每页以 N dpi 位图满幅铺底 —— 自绘 SVG/canvas 图表无法用 python-pptx 原生表达，
    位图保真 100%
  · 每页写入演讲者备注（该页 pdftotext 全文），使 PPTX 仍可检索、可复制文字
  · 页脚免责声明 / 文件编号 从备注里剔除（它们是版面固定件，不是内容）

取舍必须向用户说明：PPTX 是「位图铺底 + 备注全文」，不是原生可编辑形状。
本次交付同时给出 PDF 与 HTML —— 需要改字就改 HTML 重出，不要试图在 PPT 里改。

用法：
  python mkpptx.py <in.pdf> <out.pptx> [dpi=200] [--keep-footer]
"""
import os
import re
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.util import Pt

PT = 12700  # 1pt = 12700 EMU

# 页脚固定件特征（备注里剔除；可用 --keep-footer 关闭）
FOOTER_PAT = re.compile(
    r'(本报告基于公开行情数据|本报告版式参照贝莱德智库|数据来源：腾讯自选股|'
    r'文件编号\s|^\s*\d+\s*/\s*\d+\s*$|共\s*\d+\s*页)')


def pdfinfo(pdf):
    out = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
    pages = int(re.search(r'^Pages:\s+(\d+)', out, re.M).group(1))
    m = re.search(r'^Page size:\s+([\d.]+)\s+x\s+([\d.]+)\s+pts', out, re.M)
    if not m:
        raise SystemExit('!! 无法从 pdfinfo 读出页面尺寸')
    return pages, float(m.group(1)), float(m.group(2))


def page_text(pdf, n, keep_footer):
    t = subprocess.run(['pdftotext', '-f', str(n), '-l', str(n), '-layout', pdf, '-'],
                       capture_output=True, text=True).stdout
    lines = [re.sub(r'\s+$', '', l) for l in t.splitlines()]
    lines = [l for l in lines if l.strip()]
    if not keep_footer:
        lines = [l for l in lines if not FOOTER_PAT.search(l)]
    return '\n'.join(lines)


def render_pages(pdf, dpi, tmpdir):
    subprocess.run(['pdftoppm', '-r', str(dpi), '-png', pdf,
                    os.path.join(tmpdir, 'pg')], check=True)
    return [os.path.join(tmpdir, f)
            for f in sorted(os.listdir(tmpdir)) if f.endswith('.png')]


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    keep_footer = '--keep-footer' in sys.argv
    if len(args) < 2:
        print(__doc__)
        return 2
    src, dst = args[0], args[1]
    dpi = int(args[2]) if len(args) > 2 else 200

    n, w_pt, h_pt = pdfinfo(src)
    prs = Presentation()
    prs.slide_width = int(round(w_pt * PT))
    prs.slide_height = int(round(h_pt * PT))
    blank = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory() as td:
        imgs = render_pages(src, dpi, td)
        if len(imgs) != n:
            raise SystemExit('!! 渲染出 %d 页，PDF 有 %d 页' % (len(imgs), n))
        for i, img in enumerate(imgs, start=1):
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(img, 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
            txt = page_text(src, i, keep_footer)
            if txt:
                tf = s.notes_slide.notes_text_frame
                tf.text = txt
                for p in tf.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(11)

    prs.save(dst)
    print('PPTX 已写出: %s' % dst)
    print('  %d 页 · %.2f×%.2f pt (%.2f×%.2f in) · %d bytes'
          % (n, w_pt, h_pt, w_pt / 72, h_pt / 72, os.path.getsize(dst)))
    return 0


if __name__ == '__main__':
    sys.exit(main())
