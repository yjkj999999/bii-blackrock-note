#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
check_br.py —— 贝莱德风格 BII 简报的交付前体检（7 门）。

为什么必须跑：
  「风格重定向」最容易出的两类事故是——
   (a) **内容被悄悄改动**：正则注入误伤了正文/图表脚本。本脚本用「可见文本词元多重集」
       比对源件与产物，只允许产物多出「本次明确注入的那几句」，其余必须逐字一致。
   (b) **版式被装饰挤爆**：某页溢出成两页。本脚本核对 PDF 页数 == section 数，
       并单独核对 MediaBox（A4 横向 841.92×594.96pt）。

用法：
  python3 check_br.py OUT.html --src SRC.html --pdf OUT.pdf [--pptx OUT.pptx] \
      [--pages 22] [--json report.json]
退出码：0 全过；1 有门不过。
"""
import argparse
import collections
import io
import json
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
TOKENS_PATH = os.path.normpath(os.path.join(HERE, '..', 'assets', 'br_tokens.json'))

SCRIPT_RE = re.compile(r'<script[\s\S]*?</script>', re.I)
STYLE_RE = re.compile(r'<style[\s\S]*?</style>', re.I)
TAG_RE = re.compile(r'<[^>]+>')
CJK = re.compile(r'[\u4e00-\u9fff]')
WORD = re.compile(r'[\u4e00-\u9fff]|[A-Za-z0-9][A-Za-z0-9.,%+\-]*')


def visible_tokens(html):
    """去掉 script/style/标签，压空白，按「汉字单字 + 拉丁数字串」切词元。"""
    h = SCRIPT_RE.sub(' ', html)
    h = STYLE_RE.sub(' ', h)
    h = TAG_RE.sub(' ', h)
    h = h.replace('&nbsp;', ' ').replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>')
    return collections.Counter(WORD.findall(h))


class Checker:
    def __init__(self):
        self.rows = []
        self.ok = True

    def gate(self, no, name, passed, detail):
        self.rows.append({'gate': no, 'name': name, 'pass': bool(passed), 'detail': detail})
        if not passed:
            self.ok = False
        print('%s 门%-2s %-22s %s' % ('✅' if passed else '❌', no, name, detail))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('html')
    ap.add_argument('--src', required=True)
    ap.add_argument('--pdf', default=None)
    ap.add_argument('--pptx', default=None)
    ap.add_argument('--pages', type=int, default=None)
    ap.add_argument('--json', default=None)
    a = ap.parse_args()

    tokens = json.load(io.open(TOKENS_PATH, encoding='utf-8'))
    expected = a.pages or tokens['geometry']['expected_pages']
    C = Checker()

    out = io.open(a.html, encoding='utf-8').read()
    src = io.open(a.src, encoding='utf-8').read()

    # ── 门 1 · 结构：section 数 + 装饰件齐备 ────────────────────────────────
    n_sec = len(re.findall(r'<section\b[^>]*class="(?:cover|page)"', out))
    n_tab = len(re.findall(r'data-br-tab="', out))
    n_code = len(re.findall(r'class="br-code"', out))
    C.gate(1, '结构完整性',
           n_sec == expected and n_tab == expected and n_code == expected,
           'section %d / 书签 %d / 页脚编号 %d（期望各 %d）' % (n_sec, n_tab, n_code, expected))

    # ── 门 2 · 品牌令牌与必备装饰结构 ──────────────────────────────────────
    missing = [k for k in tokens['required_structures'] if k not in out]
    C.gate(2, '品牌令牌/装饰结构', not missing,
           '全部到位' if not missing else '缺：%s' % missing)

    # ── 门 3 · 源色无残留（色板确实被重定向了） ─────────────────────────────
    residual = {}
    for bad in tokens['forbidden_after_inject']:
        n = len(re.findall(re.escape(bad), out, re.I))
        if n:
            residual[bad] = n
    C.gate(3, '源色板无残留', not residual,
           '已清空' if not residual else '残留：%s' % residual)

    # ── 门 4 · 分类色相未塌缩（7 槽仍在） ──────────────────────────────────
    ramp = tokens['categorical_ramp']
    hits = {k: len(re.findall(re.escape(v), out, re.I)) for k, v in ramp.items() if not k.startswith('_')}
    gone = [k for k, n in hits.items() if n == 0]
    C.gate(4, '分类色相可区分', not gone,
           '7 槽齐备 %s' % hits if not gone else '丢失色相：%s' % gone)

    # ── 门 5 · 内容零改动（可见文本词元比对） ──────────────────────────────
    # 白名单不写死，而是「从 DOM 里反推本次真正注入了什么」——
    # 写死白名单只会掩盖真正的注入事故（v1.0.0 就是这么误报的）。
    ts, to = visible_tokens(src), visible_tokens(out)

    def cv_brand(h):
        m = re.search(r'<div class="cv-brand">([^<]*)<span>(.*?)</span>', h, re.S)
        return (m.group(1).strip(), m.group(2).strip()) if m else ('', '')

    s_line, s_sub = cv_brand(src)
    o_line, o_sub = cv_brand(out)
    note = re.search(r'<span class="br-style-note">(.*?)</span>', out, re.S)
    note = note.group(1) if note else ''
    code = re.search(r'<span class="br-code">(.*?)</span>', out, re.S)
    code = code.group(1) if code else ''
    n_codes = len(re.findall(r'class="br-code"', out))
    prefix = o_sub[:-len(s_sub)] if (s_sub and o_sub.endswith(s_sub)) else ''

    allowed_extra = visible_tokens(o_line) + visible_tokens(prefix) + visible_tokens(note)
    for _ in range(n_codes):
        allowed_extra += visible_tokens(code)
    # 允许消失的：仅限被替换掉的原品牌行
    allowed_gone = visible_tokens(s_line) - visible_tokens(o_line)

    only_out = to - ts
    only_src = ts - to
    bad_extra = {k: v for k, v in only_out.items() if v > allowed_extra.get(k, 0)}
    bad_gone = {k: v for k, v in only_src.items() if v > allowed_gone.get(k, 0)}
    C.gate(5, '内容零改动',
           not bad_gone and not bad_extra,
           ('可见文本与源件逐字一致（产物多出的 %d 个词元全部来自注入文案：品牌行/风格声明/页脚编号）'
            % sum(only_out.values())) if not (bad_gone or bad_extra)
           else '疑似丢失：%s ｜ 疑似误注入：%s'
                % (dict(list(bad_gone.items())[:8]), dict(list(bad_extra.items())[:8])))

    # ── 门 6 · PDF 页数与页面尺寸 ─────────────────────────────────────────
    if a.pdf and os.path.exists(a.pdf):
        info = subprocess.run(['pdfinfo', a.pdf], capture_output=True, text=True).stdout
        npg = int(re.search(r'^Pages:\s+(\d+)', info, re.M).group(1))
        m = re.search(r'^Page size:\s+([\d.]+)\s+x\s+([\d.]+)', info, re.M)
        w, h = float(m.group(1)), float(m.group(2))
        ew, eh = tokens['geometry']['media_box_pt'][2], tokens['geometry']['media_box_pt'][3]
        size_ok = abs(w - ew) < 0.6 and abs(h - eh) < 0.6
        C.gate(6, 'PDF 页数/尺寸',
               npg == expected and size_ok,
               '%d 页（期望 %d）· %.2f×%.2f pt%s'
               % (npg, expected, w, h, '' if size_ok else ' —— 尺寸不符，检查 @page'))
    else:
        C.gate(6, 'PDF 页数/尺寸', False, '未提供 --pdf 或文件不存在')

    # ── 门 7 · PPTX 结构 ─────────────────────────────────────────────────
    if a.pptx and os.path.exists(a.pptx):
        try:
            from pptx import Presentation
            prs = Presentation(a.pptx)
            ns = len(prs.slides)
            w_in = prs.slide_width / 914400.0
            h_in = prs.slide_height / 914400.0
            ew_in, eh_in = tokens['geometry']['media_box_pt'][2] / 72, tokens['geometry']['media_box_pt'][3] / 72
            pics = [sum(1 for sh in s.shapes if sh.shape_type == 13) for s in prs.slides]
            notes = [len((s.notes_slide.notes_text_frame.text or '').strip()) for s in prs.slides]
            ok = (ns == expected and all(p == 1 for p in pics) and min(notes) > 0
                  and abs(w_in - ew_in) < 0.02 and abs(h_in - eh_in) < 0.02)
            C.gate(7, 'PPTX 结构', ok,
                   '%d 页 · %.2f×%.2f in · 每页图 %s · 备注最少 %d 字'
                   % (ns, w_in, h_in, '1 张' if all(p == 1 for p in pics) else pics, min(notes)))
        except Exception as e:
            C.gate(7, 'PPTX 结构', False, '读取失败：%s' % e)
    else:
        C.gate(7, 'PPTX 结构', False, '未提供 --pptx 或文件不存在')

    print('\n%s  7 门体检：%d/7 通过'
          % ('🎉 全部通过' if C.ok else '⚠️  有门未过',
             sum(1 for r in C.rows if r['pass'])))
    if a.json:
        io.open(a.json, 'w', encoding='utf-8').write(
            json.dumps({'ok': C.ok, 'gates': C.rows}, ensure_ascii=False, indent=2))
    return 0 if C.ok else 1


if __name__ == '__main__':
    sys.exit(main())
